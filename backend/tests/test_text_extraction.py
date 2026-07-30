import io

import docx
import pytest

from app.services.text_extraction import (
    TextExtractionError,
    extract_docx,
    extract_markdown_or_text,
    extract_pdf,
    extract_text,
)
from tests.pdf_fixture import build_minimal_pdf


def test_extract_markdown_normalises_line_endings_and_bom():
    raw = "﻿heading\r\nline two\rline three\n".encode()
    text = extract_markdown_or_text(raw)
    assert text == "heading\nline two\nline three\n"


def test_extract_markdown_rejects_invalid_utf8():
    with pytest.raises(TextExtractionError):
        extract_markdown_or_text(b"\xff\xfe\x00\x01not utf8")


def test_extract_pdf_reads_real_text():
    pdf_bytes = build_minimal_pdf("Hello PDF")
    text = extract_pdf(pdf_bytes)
    assert "Hello PDF" in text


def test_extract_pdf_rejects_garbage():
    with pytest.raises(TextExtractionError):
        extract_pdf(b"not a pdf at all")


def test_extract_docx_reads_paragraphs():
    buffer = io.BytesIO()
    document = docx.Document()
    document.add_paragraph("First paragraph")
    document.add_paragraph("Second paragraph")
    document.save(buffer)

    text = extract_docx(buffer.getvalue())
    assert "First paragraph" in text
    assert "Second paragraph" in text


def test_extract_docx_rejects_garbage():
    with pytest.raises(TextExtractionError):
        extract_docx(b"not a docx at all")


def test_extract_docx_rejects_a_zip_bomb():
    """Security-review fix: DOCX is a zip container under the hood, and
    python-docx unzips it unconditionally -- a small, highly-compressible
    archive can expand to a huge decompressed size before the library
    ever gets a chance to report a format error. Constructed as a real
    zip (so the check itself, not just a garbage-input path, is
    exercised) with one entry that compresses at an enormous ratio."""
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * 50_000_000)

    with pytest.raises(TextExtractionError, match="rejected"):
        extract_docx(buffer.getvalue())


def test_extract_text_dispatches_by_extension():
    assert extract_text(".txt", b"hello") == "hello"
    with pytest.raises(TextExtractionError):
        extract_text(".exe", b"whatever")
