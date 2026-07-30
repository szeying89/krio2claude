import openpyxl
import pytest

from app.services.cri.workbook_parser import CRIWorkbookParseError, parse_workbook

FIXTURE = "tests/fixtures/cri/sample_cri_profile.xlsx"


def test_parses_only_diagnostic_statement_rows():
    catalog = parse_workbook(FIXTURE)
    ids = {s.profile_id for s in catalog.statements}
    assert ids == {"GV.OC-01.01", "GV.OC-01.02", "PR.AA-05.01"}


def test_csf_path_split_from_function_category_subcategory_column():
    catalog = parse_workbook(FIXTURE)
    statement = catalog.by_profile_id("GV.OC-01.01")
    assert statement.csf_path == ("GOVERN", "Organizational Context", "Organizational Mission")


def test_applicable_tiers_from_yes_columns():
    catalog = parse_workbook(FIXTURE)
    assert catalog.by_profile_id("GV.OC-01.01").applicable_tiers == (1, 2, 3, 4)
    assert catalog.by_profile_id("GV.OC-01.02").applicable_tiers == (1, 2)
    assert catalog.by_profile_id("PR.AA-05.01").applicable_tiers == (1,)


def test_regulatory_references_parsed_with_counts():
    catalog = parse_workbook(FIXTURE)
    refs = {r.short_code: r.count for r in catalog.by_profile_id("GV.OC-01.01").regulatory_references}
    assert refs == {"TESTREG-A": 2, "TESTREG-B": 1}


def test_statement_with_no_regulatory_references_is_empty_tuple():
    catalog = parse_workbook(FIXTURE)
    assert catalog.by_profile_id("PR.AA-05.01").regulatory_references == ()


def test_unresolved_regulatory_reference_is_surfaced_not_dropped():
    catalog = parse_workbook(FIXTURE)
    assert catalog.unresolved_regulatory_references == ("TESTREG-UNKNOWN",)
    # and it must still be visible on the statement itself, not silently stripped
    codes = {r.short_code for r in catalog.by_profile_id("GV.OC-01.02").regulatory_references}
    assert "TESTREG-UNKNOWN" in codes


def test_regulatory_documents_resolve_known_codes():
    catalog = parse_workbook(FIXTURE)
    doc = catalog.regulatory_documents["TESTREG-A"]
    assert doc.document_name == "Test Regulation A"
    assert doc.issuing_organization == "Test Regulatory Authority A"


def test_subject_tags_and_eee_packages_merged_from_assessment_sheet():
    catalog = parse_workbook(FIXTURE)
    statement = catalog.by_profile_id("GV.OC-01.01")
    assert statement.subject_tags == ("#architecture", "#mission_and_strategy", "#risk_management")
    assert statement.eee_package_ids == ("EEE-001", "EEE-002")
    assert statement.name == "Governance alignment"


def test_statement_with_no_tags_or_eee_defaults_to_empty():
    catalog = parse_workbook(FIXTURE)
    statement = catalog.by_profile_id("PR.AA-05.01")
    assert statement.subject_tags == ()
    assert statement.eee_package_ids == ()


def test_eee_packages_catalog_parsed():
    catalog = parse_workbook(FIXTURE)
    assert catalog.eee_packages["EEE-001"].name == "Cybersecurity and IT Strategy Documents"
    assert "Example evidence item 1" in catalog.eee_packages["EEE-001"].example_evidence


def test_missing_sheet_raises_actionable_error(tmp_path):
    wb = openpyxl.Workbook()
    wb.save(tmp_path / "broken.xlsx")
    with pytest.raises(CRIWorkbookParseError, match="CRI Profile v2.2 Structure"):
        parse_workbook(tmp_path / "broken.xlsx")


def test_missing_required_column_raises_actionable_error(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CRI Profile v2.2 Structure"
    ws["A1"] = "Outline Id"
    ws["B1"] = "Level"
    # missing the rest of the required columns
    wb.save(tmp_path / "missing_columns.xlsx")
    with pytest.raises(CRIWorkbookParseError, match="missing expected column"):
        parse_workbook(tmp_path / "missing_columns.xlsx")


def test_rejects_a_zip_bomb_workbook(tmp_path):
    """Security-review fix: XLSX is a zip container too, and openpyxl
    unzips it unconditionally -- checked (and rejected) before openpyxl
    ever gets to parse it, given a real zip with one anomalously
    compressible entry."""
    import zipfile

    archive_path = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(archive_path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"0" * 50_000_000)

    with pytest.raises(CRIWorkbookParseError, match="rejected"):
        parse_workbook(archive_path)
