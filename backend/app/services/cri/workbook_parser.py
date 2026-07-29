"""CRI Profile v2.2 workbook parser.

Verified against a real CRI Profile v2.2 workbook. Column layout is
resolved by scanning for a known header anchor within the first several
rows of each sheet (not a fixed row number) and matching whitespace-
normalized header text — a version bump that reflows rows still parses,
while a genuinely different layout (missing/renamed columns) fails with an
actionable error naming the sheet and the missing column, rather than
silently guessing.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import IO, Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.services.cri.models import (
    ControlObjectiveCatalog,
    DiagnosticStatement,
    EEEPackage,
    RegulatoryDocument,
    RegulatoryReference,
)

STRUCTURE_SHEET = "CRI Profile v2.2 Structure"
ASSESSMENT_SHEET = "CRI Profile v2.2 Assessment"
CATALOG_SHEET = "Catalog of Mapped Documents"
EEE_SHEET = "EEE Packages"

STRUCTURE_REQUIRED = [
    "Outline Id",
    "Level",
    "Profile Id",
    "CRI Profile Function / Category / Subcategory",
    "CRI Profile v2.2 Diagnostic Statement",
    "Tier-1",
    "Tier-2",
    "Tier-3",
    "Tier-4",
    "Financial Services Mapping References",
]
ASSESSMENT_REQUIRED = [
    "Profile Id",
    "Examples of Effective Evidence (EEE) Packages",
    "Profile Subject Tags",
]
CATALOG_REQUIRED = [
    "Mapping Worksheet Link",
    "Document Name",
    "Issuing Organization",
    "Region",
    "Issue Date",
    "Source Document Link",
    "Version 2.2 Status",
]
EEE_REQUIRED = ["Package Id", "EEE Package Name", "Package Example Evidence Items"]

_REG_REF_RE = re.compile(r"^(?P<code>.+?)\s*\((?P<count>\d+)\)$")
_EEE_ID_RE = re.compile(r"^\s*(EEE-\d+)")


class CRIWorkbookParseError(Exception):
    pass


def _normalize(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


def _get_sheet(wb: Any, name: str) -> Worksheet:
    if name in wb.sheetnames:
        return wb[name]
    normalized = {sheet.strip(): sheet for sheet in wb.sheetnames}
    if name in normalized:
        return wb[normalized[name]]
    raise CRIWorkbookParseError(
        f"workbook is missing the {name!r} sheet; found sheets: {wb.sheetnames!r}"
    )


def _find_header_row(ws: Worksheet, anchor: str, max_scan: int = 15) -> int:
    anchor_norm = _normalize(anchor)
    for row in ws.iter_rows(min_row=1, max_row=max_scan):
        for cell in row:
            if _normalize(cell.value) == anchor_norm:
                return int(cell.row)
    raise CRIWorkbookParseError(
        f"sheet {ws.title!r}: could not locate a header row containing {anchor!r} "
        f"in the first {max_scan} rows — unexpected layout, refusing to guess"
    )


def _header_index(ws: Worksheet, header_row: int, required: list[str]) -> dict[str, int]:
    headers: dict[str, int] = {}
    for cell in ws[header_row]:
        norm = _normalize(cell.value)
        if norm:
            headers[norm] = cell.column
    missing = [h for h in required if h not in headers]
    if missing:
        raise CRIWorkbookParseError(
            f"sheet {ws.title!r}: missing expected column(s) {missing!r} in header row "
            f"{header_row}; found {sorted(headers)!r}"
        )
    return headers


def _parse_regulatory_references(raw: Any) -> tuple[RegulatoryReference, ...]:
    if not raw:
        return ()
    refs = []
    for part in str(raw).split(","):
        part = part.strip()
        if not part:
            continue
        match = _REG_REF_RE.match(part)
        if not match:
            raise CRIWorkbookParseError(f"could not parse regulatory reference entry {part!r}")
        refs.append(
            RegulatoryReference(short_code=match.group("code").strip(), count=int(match.group("count")))
        )
    return tuple(refs)


def _parse_structure_sheet(ws: Worksheet) -> list[DiagnosticStatement]:
    header_row = _find_header_row(ws, "Outline Id")
    cols = _header_index(ws, header_row, STRUCTURE_REQUIRED)

    statements = []
    for row in ws.iter_rows(min_row=header_row + 1):
        level = _normalize(row[cols["Level"] - 1].value)
        if level != "DS":
            continue

        outline_id = _normalize(row[cols["Outline Id"] - 1].value)
        profile_id = _normalize(row[cols["Profile Id"] - 1].value)
        csf_raw = _normalize(row[cols["CRI Profile Function / Category / Subcategory"] - 1].value)
        text = _normalize(row[cols["CRI Profile v2.2 Diagnostic Statement"] - 1].value)

        tiers = tuple(
            tier
            for tier, header in ((1, "Tier-1"), (2, "Tier-2"), (3, "Tier-3"), (4, "Tier-4"))
            if _normalize(row[cols[header] - 1].value) == "Yes"
        )
        reg_refs = _parse_regulatory_references(
            row[cols["Financial Services Mapping References"] - 1].value
        )

        statements.append(
            DiagnosticStatement(
                outline_id=outline_id,
                profile_id=profile_id,
                csf_path=tuple(p.strip() for p in csf_raw.split("/") if p.strip()),
                name="",
                text=text,
                applicable_tiers=tiers,
                regulatory_references=reg_refs,
            )
        )
    return statements


def _parse_assessment_sheet(ws: Worksheet) -> dict[str, tuple[tuple[str, ...], tuple[str, ...], str]]:
    """Returns {profile_id: (subject_tags, eee_package_ids, diagnostic_statement_name)}."""
    header_row = _find_header_row(ws, "Outline Id")
    cols = _header_index(ws, header_row, [*ASSESSMENT_REQUIRED, "Diagnostic Statement Name"])

    result: dict[str, tuple[tuple[str, ...], tuple[str, ...], str]] = {}
    for row in ws.iter_rows(min_row=header_row + 1):
        profile_id = _normalize(row[cols["Profile Id"] - 1].value)
        if not profile_id:
            continue

        tags_raw = row[cols["Profile Subject Tags"] - 1].value
        tags = tuple(t.strip() for t in str(tags_raw or "").split("|") if t.strip())

        eee_raw = row[cols["Examples of Effective Evidence (EEE) Packages"] - 1].value
        eee_ids = []
        for line in str(eee_raw or "").splitlines():
            match = _EEE_ID_RE.match(line)
            if match:
                eee_ids.append(match.group(1))

        name = _normalize(row[cols["Diagnostic Statement Name"] - 1].value)
        result[profile_id] = (tags, tuple(eee_ids), name)
    return result


def _parse_catalog_sheet(ws: Worksheet) -> dict[str, RegulatoryDocument]:
    header_row = _find_header_row(ws, "Mapping Worksheet Link")
    cols = _header_index(ws, header_row, CATALOG_REQUIRED)

    documents: dict[str, RegulatoryDocument] = {}
    for row in ws.iter_rows(min_row=header_row + 1):
        short_code = _normalize(row[cols["Mapping Worksheet Link"] - 1].value)
        if not short_code:
            continue
        documents[short_code] = RegulatoryDocument(
            short_code=short_code,
            document_name=_normalize(row[cols["Document Name"] - 1].value),
            issuing_organization=_normalize(row[cols["Issuing Organization"] - 1].value),
            region=_normalize(row[cols["Region"] - 1].value),
            issue_date=_normalize(row[cols["Issue Date"] - 1].value),
            source_link=_normalize(row[cols["Source Document Link"] - 1].value),
            status=_normalize(row[cols["Version 2.2 Status"] - 1].value),
        )
    return documents


def _parse_eee_sheet(ws: Worksheet) -> dict[str, EEEPackage]:
    header_row = _find_header_row(ws, "Package Id")
    cols = _header_index(ws, header_row, EEE_REQUIRED)

    packages: dict[str, EEEPackage] = {}
    for row in ws.iter_rows(min_row=header_row + 1):
        package_id = _normalize(row[cols["Package Id"] - 1].value)
        if not package_id:
            continue
        packages[package_id] = EEEPackage(
            id=package_id,
            name=_normalize(row[cols["EEE Package Name"] - 1].value),
            example_evidence=_normalize(row[cols["Package Example Evidence Items"] - 1].value),
        )
    return packages


def parse_workbook(source: str | Path | IO[bytes], version: str = "2.2") -> ControlObjectiveCatalog:
    wb = load_workbook(source, data_only=True, read_only=True)

    structure_statements = _parse_structure_sheet(_get_sheet(wb, STRUCTURE_SHEET))
    assessment_index = _parse_assessment_sheet(_get_sheet(wb, ASSESSMENT_SHEET))
    regulatory_documents = _parse_catalog_sheet(_get_sheet(wb, CATALOG_SHEET))
    eee_packages = _parse_eee_sheet(_get_sheet(wb, EEE_SHEET))

    merged_statements = []
    for statement in structure_statements:
        tags, eee_ids, name = assessment_index.get(statement.profile_id, ((), (), ""))
        merged_statements.append(
            DiagnosticStatement(
                outline_id=statement.outline_id,
                profile_id=statement.profile_id,
                csf_path=statement.csf_path,
                name=name,
                text=statement.text,
                applicable_tiers=statement.applicable_tiers,
                regulatory_references=statement.regulatory_references,
                subject_tags=tags,
                eee_package_ids=eee_ids,
            )
        )

    referenced_codes = {
        ref.short_code for s in merged_statements for ref in s.regulatory_references
    }
    unresolved = tuple(sorted(referenced_codes - set(regulatory_documents)))

    return ControlObjectiveCatalog(
        version=version,
        statements=tuple(merged_statements),
        regulatory_documents=regulatory_documents,
        eee_packages=eee_packages,
        unresolved_regulatory_references=unresolved,
    )
